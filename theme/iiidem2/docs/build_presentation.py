#!/usr/bin/env python3
"""Build IIIDEM Moodle LMS project PowerPoint from screenshots."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ASSETS = Path(__file__).resolve().parent / "ppt_assets"
OUTPUT = Path(__file__).resolve().parent / "IIIDEM_Moodle_LMS_Project.pptx"

NAVY = RGBColor(0, 66, 116)
GOLD = RGBColor(232, 163, 23)
WHITE = RGBColor(255, 255, 255)
DARK = RGBColor(10, 37, 64)


def set_title(slide, title: str, subtitle: str = "") -> None:
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle
    for shape in [slide.shapes.title, slide.placeholders[1]]:
        for p in shape.text_frame.paragraphs:
            for run in p.runs:
                run.font.name = "Calibri"
    slide.shapes.title.text_frame.paragraphs[0].runs[0].font.color.rgb = NAVY
    slide.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(32)
    if subtitle:
        slide.placeholders[1].text_frame.paragraphs[0].runs[0].font.size = Pt(18)
        slide.placeholders[1].text_frame.paragraphs[0].runs[0].font.color.rgb = DARK


def add_bullets(slide, items: list[str]) -> None:
    body = slide.placeholders[1].text_frame
    body.clear()
    for i, item in enumerate(items):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(20)
        p.font.name = "Calibri"


def add_image_slide(prs: Presentation, title: str, image_name: str, caption: str) -> None:
    layout = prs.slide_layouts[5]  # Title only
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title
    slide.shapes.title.text_frame.paragraphs[0].runs[0].font.color.rgb = NAVY
    slide.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(28)

    img_path = ASSETS / image_name
    if img_path.exists():
        slide.shapes.add_picture(str(img_path), Inches(0.4), Inches(1.2), width=Inches(9.2))

    box = slide.shapes.add_textbox(Inches(0.4), Inches(6.85), Inches(9.2), Inches(0.6))
    tf = box.text_frame
    tf.text = caption
    tf.paragraphs[0].font.size = Pt(14)
    tf.paragraphs[0].font.color.rgb = DARK


def main() -> None:
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # Title slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "IIIDEM Certification LMS"
    slide.placeholders[1].text = (
        "Moodle-based learning platform\n"
        "Live Webex classes · MCQ assessments · Teacher dashboard\n"
        "iiidem-certification.ddev.site"
    )
    slide.shapes.title.text_frame.paragraphs[0].runs[0].font.color.rgb = NAVY
    slide.shapes.title.text_frame.paragraphs[0].runs[0].font.size = Pt(40)

    # Agenda
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "Project overview", "Custom Moodle theme + plugins for IIIDEM")
    add_bullets(slide, [
        "Public marketing site & course catalogue",
        "Student & teacher dashboards (theme_iiidem2)",
        "Live Webex virtual classroom pages",
        "Live class MCQ during sessions (local_iiidem_livequiz)",
        "Quiz assessments & teacher reporting",
        "Payment gateway (PNB) for course enrolment",
    ])

    # Platform
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "Technology stack", "")
    add_bullets(slide, [
        "Moodle 4.5 LMS (core platform)",
        "Custom theme: theme/iiidem2",
        "Local plugins: coursereviews, iiidem_livequiz, PNB payment gateway",
        "Activities: Page (Webex), Quiz, Attendance, Certificate",
        "DDEV local environment for development",
    ])

    # Screenshots
    screenshots = [
        ("Home page", "01-home.png", "Marketing front page with course catalogue and IIIDEM branding."),
        ("Login", "02-login.png", "Secure login for students, teachers, and administrators."),
        ("Course page", "07-course-page.png", "Certificate course: AI, Elections & Democratic Governance (POB101)."),
        ("Teacher dashboard", "04-teacher-dashboard.png", "Instructor console: live class control, roster, assessments, attendance."),
        ("Live Webex class", "05-live-class.png", "Students join Webex while keeping the Moodle live class page open."),
        ("Live class MCQ", "06-live-mcq-manage.png", "Teachers create & release MCQ questions during live sessions."),
        ("Week 1 quiz", "08-quiz-assessment.png", "Structured Moodle quiz for weekend assessments."),
    ]

    for title, img, caption in screenshots:
        add_image_slide(prs, title, img, caption)

    # Workflow slide
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "Live class + MCQ workflow", "Option 3 — implemented")
    add_bullets(slide, [
        "1. Teacher opens Dashboard → Live class MCQ → Manage",
        "2. Add questions → Release to students",
        "3. Students open Live class page → answer in Live questions panel",
        "4. Teacher sees submissions & scores in real time",
        "5. Video/audio stays in Webex; answers stored in Moodle",
    ])

    # Features slide
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    set_title(slide, "Key features delivered", "")
    add_bullets(slide, [
        "Custom teacher dashboard with assessment summary & per-student quiz tracking",
        "Live virtual classroom UI with Join Webex button",
        "Full-screen custom MCQ quiz attempt experience",
        "Student course reviews & testimonials",
        "Course fee payment via PNB gateway",
        "Attendance & certificate management",
    ])

    # Closing
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Thank you"
    slide.placeholders[1].text = "IIIDEM · Election Commission of India\nMoodle LMS Project Demo"

    prs.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    main()
