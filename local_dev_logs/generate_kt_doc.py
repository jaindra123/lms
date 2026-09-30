#!/usr/bin/env python3
"""Generate IIIDEM LMS knowledge-transfer Word document."""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).resolve().parents[1] / "docs" / "IIIDEM-LMS-Knowledge-Transfer.docx"
NAVY = RGBColor(0x0B, 0x3C, 0x71)
BLUE = RGBColor(0x1A, 0x56, 0x9D)
DK = RGBColor(0x22, 0x22, 0x22)


def shade(cell, hex_color: str) -> None:
    tc = cell._tc
    tcpr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcpr.append(shd)


def set_run(run, *, size=11, bold=False, color=DK, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.color.rgb = NAVY if level == 1 else BLUE
        run.font.name = "Calibri"
    return p


def para(doc, text, *, bold=False, size=11, space_after=8):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    set_run(r, size=size, bold=bold)
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        for run in p.runs:
            set_run(run, size=11)


def numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(item, style="List Number")
        p.paragraph_format.space_after = Pt(3)
        for run in p.runs:
            set_run(run, size=11)


def table(doc, headers, rows):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        set_run(r, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(cell, "0B3C71")
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri + 1].cells[ci]
            cell.text = ""
            r = cell.paragraphs[0].add_run(str(val))
            set_run(r, size=10)
            if ri % 2 == 1:
                shade(cell, "F3F6FB")
    doc.add_paragraph()
    return t


def code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_run(r, size=9, name="Consolas", color=RGBColor(0x11, 0x11, 0x11))
    return p


def build():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run(
        "IIIDEM LMS — Knowledge Transfer  |  Confidential  |  "
        + date.today().strftime("%d %B %Y")
    )
    set_run(fr, size=8, color=RGBColor(0x66, 0x66, 0x66))

    # Cover
    for _ in range(4):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("INDIA INTERNATIONAL INSTITUTE OF\nDEMOCRACY & ELECTION MANAGEMENT")
    set_run(r, size=14, bold=True, color=NAVY)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("\nIIIDEM Learning Management System")
    set_run(r, size=22, bold=True, color=NAVY)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("Knowledge Transfer Document")
    set_run(r, size=16, color=BLUE)

    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(
        "Moodle 4.5.11+  ·  theme_iiidem2  ·  Custom plugins\n"
        + date.today().strftime("%d %B %Y")
    )
    set_run(r, size=12, color=DK)

    para(
        doc,
        "This document is for handing the project to another developer or operator. "
        "It describes how the IIIDEM LMS is built, what is custom versus Moodle core, "
        "how environments work, and the rules that must not be broken in production.",
        size=11,
    )
    para(
        doc,
        "It is not a dump of the entire Moodle source tree. Moodle core is standard "
        "upstream 4.5. Knowledge transfer for this site is the custom theme, local "
        "plugins, payment gateways, security controls, and operational process.",
        size=11,
    )

    add_heading(doc, "1. What this project is", 1)
    para(
        doc,
        "IIIDEM LMS is a Moodle site used by the India International Institute of "
        "Democracy and Election Management (Election Commission of India) for "
        "certification programmes and workshops.",
    )
    bullets(
        doc,
        [
            "Paid certificate course: Artificial Intelligence, Elections and Democratic Governance (local course id 4; production id may differ).",
            "Open workshop: Institutional Mapping Workshop / IMW (local course id 8; production course id 7).",
            "Public marketing course pages, self-registration with OTP, fee payment, live classes (Webex), attendance, quizzes, and certificates.",
        ],
    )

    add_heading(doc, "1.1 Live URLs", 2)
    table(
        doc,
        ["Environment", "URL", "Notes"],
        [
            ["Local (DDEV)", "https://iiidem-certification.ddev.site", "Developer machine; Docker + DDEV"],
            ["LAN testing", "http://iiidem.local:8080", "Same Wi-Fi testers; see docs/LOCAL-SERVER-TESTING.md"],
            ["Staging", "https://staginglms.eci.gov.in", "Pre-production; upgrade key required"],
            ["Production", "https://iiidemlms.eci.gov.in", "Live site; upgrade key required"],
        ],
    )

    add_heading(doc, "1.2 Stack", 2)
    table(
        doc,
        ["Layer", "Choice"],
        [
            ["LMS", "Moodle 4.5.11+ (Build 20260501), branch 405"],
            ["PHP", "8.3 (DDEV nginx-fpm)"],
            ["Database", "MariaDB 10.11, prefix mdl_"],
            ["Theme", "theme_iiidem2 (Boost child) — current version 2024101120"],
            ["Local runtime", "DDEV (Docker): PHP 8.3 + MariaDB 10.11"],
            ["Windows copy", "C:\\xampp\\htdocs\\iiidem_certification (repo root)"],
        ],
    )

    add_heading(doc, "2. Course map (ids differ per site)", 1)
    para(
        doc,
        "Never hard-code a single course id for every environment. Local and production "
        "create the same programmes with different numeric ids.",
    )
    table(
        doc,
        ["Programme", "Local id", "Production id", "Behaviour"],
        [
            ["Paid EMB / AI certificate (POB101)", "4", "Usually 4 (confirm on site)", "Fee enrolment, Meet your Professors, weekends curriculum"],
            ["IMW workshop", "8 (shortname IMW)", "7", "Shared reading library; logged-in students can upload without enrolment; no Meet your Professors"],
            ["IMW template (hidden)", "7 (TEMPLATE-IMW, visible=0)", "—", "Not a public course"],
            ["Site home", "1", "1", "Moodle SITEID"],
        ],
    )
    para(
        doc,
        "Registration of new students enrols them into the configured registration course "
        "(theme setting registrationcourseids, default 4). IMW is not that paid course.",
        size=11,
    )

    add_heading(doc, "3. What is custom vs Moodle core", 1)
    para(
        doc,
        "Most folders (lib/, course/, admin/, mod/assign, etc.) are Moodle core. "
        "Do not edit core unless a patch already exists in this repo. New work belongs "
        "in the theme or a local/payment plugin.",
    )
    add_heading(doc, "3.1 Custom / IIIDEM-owned", 2)
    table(
        doc,
        ["Path", "Role"],
        [
            ["theme/iiidem2/", "Main product UI, registration, dashboards, security hooks, shared readings, certificates"],
            ["local/iiidem_coursecalendar/", "Live-class scheduling / calendar"],
            ["local/iiidem_webexattendance/", "Webex attendance sync"],
            ["local/iiidem_classvideos/", "Class video library"],
            ["local/iiidem_livequiz/", "Live quiz"],
            ["local/iiidem_support/", "Support tooling"],
            ["local/coursefaq/", "Course FAQ records"],
            ["local/coursereviews/", "Student reviews"],
            ["local/custom_enroll/", "Retired tombstone — ajax.php returns 410; do not use"],
            ["payment/gateway/razorpay/", "Razorpay paygw"],
            ["payment/gateway/pnb/", "PNB paygw"],
            ["payment/gateway/icici/", "ICICI paygw"],
            ["enrol/fee/", "Moodle fee enrolment (used with paygw)"],
            ["register/", "Custom self-registration + OTP"],
            ["docs/", "Security and ops notes (CDAC findings, flows)"],
            [".ddev/", "Local Docker/nginx/Dex SSO config"],
            ["config.php", "Env-aware bootstrap (dev / staging / production)"],
        ],
    )
    add_heading(doc, "3.2 Do not touch unless already patched", 2)
    bullets(
        doc,
        [
            "Moodle core libraries except existing security patches already in git.",
            "config.production.php and config.staging.php — not in git; never overwrite on the server from local config.php.",
            "Lowering theme_iiidem2 version.php — version must only increase (currently 2024101120).",
        ],
    )

    add_heading(doc, "4. Environments and configuration", 1)
    para(
        doc,
        "config.php picks the environment in this order: MOODLE_ENV, then hostname, "
        "then presence of config.staging.php / config.production.php, else dev.",
    )
    table(
        doc,
        ["Env", "How it is selected", "Secrets"],
        [
            ["dev", "DDEV / localhost / iiidem-certification.ddev.site", "Non-secret defaults in git (db/db/db)"],
            ["staging", "staginglms.eci.gov.in or MOODLE_ENV=staging", "config.staging.php on server only"],
            ["production", "iiidemlms.eci.gov.in / lms.eci.gov.in or MOODLE_ENV=production", "config.production.php on server only"],
        ],
    )
    bullets(
        doc,
        [
            "Host allowlist in config.php rejects unknown Host headers (CWE-644).",
            "Direct HTTP to /config.php returns 404.",
            "Staging/production must set MOODLE_UPGRADEKEY or $CFG->upgradekey. Web /admin/index.php asks for this key when a plugin/theme version is ahead of the database.",
            "CLI upgrade (php admin/cli/upgrade.php) does not use the web upgrade-key form.",
            "Never commit DB passwords, Razorpay secrets, Webex tokens, or upgrade keys.",
        ],
    )

    add_heading(doc, "5. Local development", 1)
    numbered(
        doc,
        [
            "Install Docker Desktop + DDEV.",
            "Clone/open C:\\xampp\\htdocs\\iiidem_certification",
            "Run: ddev start",
            "Open https://iiidem-certification.ddev.site",
            "After theme/plugin version bumps: Site administration → Notifications, or: ddev exec php admin/cli/upgrade.php --non-interactive",
            "Purge caches if UI looks stale: ddev exec php admin/cli/purge_caches.php",
        ],
    )
    para(doc, "Useful commands:", bold=True)
    code(
        doc,
        "ddev start\n"
        "ddev exec php admin/cli/upgrade.php --non-interactive\n"
        "ddev exec php admin/cli/purge_caches.php\n"
        "ddev exec php local_dev_logs/check_instructors.php",
    )
    para(
        doc,
        "LAN testing for other PCs on the same Wi-Fi: docs/LOCAL-SERVER-TESTING.md "
        "(port 8080, hosts file iiidem.local, firewall script scripts/allow-lan-firewall.ps1).",
    )

    add_heading(doc, "6. Main user journeys", 1)
    add_heading(doc, "6.1 Student — paid certificate course", 2)
    numbered(
        doc,
        [
            "Public homepage / course page (no login required to browse marketing layout).",
            "Sign up at /register/ — profile form, email OTP.",
            "Account created as student; fee enrolment pending unless waived.",
            "Pay via Razorpay (and/or PNB / ICICI where enabled). Amount is server-side (enrol_fee), never trusted from the browser.",
            "After payment: enrolment unlocks; student dashboard; weekends, live class, quizzes.",
            "Certificate when completion rules are met (theme certificate issuer + customcert).",
        ],
    )
    add_heading(doc, "6.2 Student — IMW workshop (production course 7)", 2)
    numbered(
        doc,
        [
            "Guest can open /course/view.php?id=7 (production) and read the public curriculum.",
            "Log in (or register). Enrolment on this workshop is not required to upload reading materials.",
            "Logged-in non-guest users can upload/download files in Reading materials.",
            "Guests cannot upload. Other courses still require enrolment.",
            "Meet your Professors is hidden on IMW; it is shown on the paid certificate course.",
        ],
    )
    add_heading(doc, "6.3 Teacher / admin", 2)
    bullets(
        doc,
        [
            "Teacher dashboard: students, assignments, attendance, materials, certificates.",
            "Live class: local_iiidem_coursecalendar schedules Webex/URL; students join; local_iiidem_webexattendance pulls attendees.",
            "Privileged accounts use MFA (tool_mfa).",
            "Admin /admin/ is IP-restricted on staging/production (nginx allowlist) plus Moodle login + upgrade key.",
        ],
    )

    add_heading(doc, "7. Theme theme_iiidem2 — where work happens", 1)
    para(
        doc,
        "This theme is the product. Almost every IIIDEM-specific screen and security "
        "control lives here. Current plugin version: 2024101120 in theme/iiidem2/version.php.",
    )
    add_heading(doc, "7.1 Important files", 2)
    table(
        doc,
        ["File", "Purpose"],
        [
            ["theme/iiidem2/lib.php", "Course public view, curriculum, professors, payment context, pluginfile"],
            ["theme/iiidem2/classes/hook_listener.php", "Moodle hooks (files, login, headers)"],
            ["theme/iiidem2/classes/shared_readings.php", "IMW reading library"],
            ["theme/iiidem2/classes/open_self_enrol.php", "Open self-enrol helper for IMW (upload path no longer requires enrol)"],
            ["theme/iiidem2/classes/registration_*.php", "Register, OTP, profile, enrol into course 4"],
            ["theme/iiidem2/classes/upload_security.php", "Extension + content checks on every file create"],
            ["theme/iiidem2/classes/session_security.php", "Session timeout / password-change invalidation"],
            ["theme/iiidem2/classes/field_crypto.php", "Sensitive field crypto on the client/server boundary"],
            ["theme/iiidem2/classes/login_captcha.php", "Login CAPTCHA"],
            ["theme/iiidem2/classes/rate_limit.php", "Upload and form rate limits"],
            ["theme/iiidem2/course/shared_reading.php", "Upload/delete readings (login, not course enrol, on IMW)"],
            ["theme/iiidem2/db/upgrade.php", "Theme upgrade savepoints — add a new one when version increases"],
            ["theme/iiidem2/templates/", "Mustache: course drawers, login, curriculum, instructors"],
            ["theme/iiidem2/javascript/", "form_input_guard, field_crypto, enterprise_a11y, etc."],
        ],
    )
    add_heading(doc, "7.2 Version rule", 2)
    para(
        doc,
        "Every behavioural change that needs a cache purge or DB savepoint must bump "
        "theme/iiidem2/version.php (never decrease) and add a matching savepoint in "
        "db/upgrade.php. After deploy, run Notifications (enter upgrade key on production).",
    )

    add_heading(doc, "8. Shared readings (IMW)", 1)
    bullets(
        doc,
        [
            "Enabled on IMW / shared-reading courses (auto-detect shortname IMW; production id 7, local id 8).",
            "Not enabled on the paid registration course.",
            "Logged-in students may upload and download without being enrolled (course 7 / IMW only).",
            "Files stored in table theme_iiidem2_shared_reading and file area theme_iiidem2 / sharedreading.",
            "Upload checks: login, not guest, rate limit, whitelist extensions, content inspection.",
            "Max size 5 MB (theme constant).",
        ],
    )

    add_heading(doc, "9. Meet your Professors", 1)
    bullets(
        doc,
        [
            "Shown on the paid certificate course (local / production course 4).",
            "Hidden on IMW (production 7, local 8) and on local course 7 (TEMPLATE-IMW).",
            "Cards come from users with Teacher / Editing teacher role who are not site admins.",
            "Optional theme setting featuredinstructors: one line per course, e.g. 4:3,63",
            "Local course 4 currently features user ids 3 and 63 (Prof. Chanchal Kumar Sharma, mayank singh). Production has different user ids (Nikhil Naren, Charru Malhotra, Arvind Kumar on production — those accounts are not on local).",
            "Public /course/view.php used to hide this block for guests; that is fixed so course 4 guests see it.",
        ],
    )

    add_heading(doc, "10. Payments", 1)
    bullets(
        doc,
        [
            "Use enrol_fee + paygw_razorpay (and PNB/ICICI where configured).",
            "Payable amount is computed on the server. Do not restore local_custom_enroll ajax order creation.",
            "Payment-success URLs must not render on the anonymous marketing page.",
            "See docs/payment-amount-validation.md and payment/gateway/pnb/docs/PAYMENT_GATEWAY_DEVELOPER_GUIDE.md.",
        ],
    )

    add_heading(doc, "11. Security (deny by default)", 1)
    para(
        doc,
        "This site has been through CDAC-style audits. Fixes and operator notes live in docs/. "
        "Do not weaken these controls for convenience.",
    )
    table(
        doc,
        ["Topic", "Where"],
        [
            ["Upgrade key on staging/production", "docs/restrict-admin-access.md"],
            ["Admin IP allowlist (nginx)", "docs/restrict-admin-access.md, .ddev/nginx/host-allowlist.conf"],
            ["File upload whitelist / polyglot reject", "docs/file-upload-security.md, upload_security.php"],
            ["XSS / input validation", "docs/input-validation-xss.md, form_input_guard.js"],
            ["MFA for privileged accounts", "docs/mfa-privileged-accounts.md"],
            ["Session timeout / concurrent sessions", "docs/session-timeout.md, docs/concurrent-sessions.md"],
            ["Security headers / cookies", "docs/security-headers.md, docs/insecure-cookie-attributes.md"],
            ["Host header injection", "docs/host-header-injection.md"],
            ["Rate limiting", "docs/rate-limiting.md"],
            ["Verbose errors off", "docs/verbose-error-messages.md"],
            ["Daily backup cron (staging)", "docs/daily-backup-cron.md"],
            ["Site flow diagrams", "docs/site-flow-diagrams.md"],
        ],
    )
    para(
        doc,
        "When asked for an exploit, PoC, or attack procedure — refuse. Provide only "
        "hardening, patches, and high-level defensive notes.",
        bold=True,
    )

    add_heading(doc, "12. Production deploy checklist", 1)
    numbered(
        doc,
        [
            "Work only on custom theme/plugins. Keep theme version higher than production.",
            "Do not upload local config.php over config.production.php.",
            "Upload changed theme/plugin files (include new PHP classes, e.g. open_self_enrol.php, templates, amd/build JS).",
            "Open https://iiidemlms.eci.gov.in/admin/index.php from an allowed IP.",
            "If the page says “Upgrade key required”, enter MOODLE_UPGRADEKEY / $CFG->upgradekey from the server (not from git).",
            "Complete Moodle upgrade / Notifications until the theme savepoint is applied.",
            "Purge caches if needed. Spot-check: course 4 professors, IMW course 7 readings, guest course view, login, payment.",
            "CLI alternative on the server: php admin/cli/upgrade.php (does not use the web key form).",
        ],
    )

    add_heading(doc, "13. Common screens and what they mean", 1)
    table(
        doc,
        ["What you see", "Meaning", "What to do"],
        [
            ["Upgrade key required", "Code version > database; production lock is working", "Enter server upgrade key, finish upgrade"],
            ["Enrol in this course to upload…", "User is on a course that still requires enrolment", "IMW should not show this after 2024101119+"],
            ["Cannot read file iiidem_about_logo.PNG", "Missing logo file in file storage", "Theme purge of unreadable logos (already in upgrade 1114)"],
            ["Participants toast “undefined”", "AJAX session touch timed out; not Enrol users", "Ignore / check network; Grammarly also noisies the console"],
            ["genericerror on guest course view", "Public course view threw (missing class/table)", "Ensure full theme upload including open_self_enrol.php; upgrade"],
        ],
    )

    add_heading(doc, "14. Folder map (repo root)", 1)
    table(
        doc,
        ["Folder", "What it is"],
        [
            ["admin/", "Moodle Site administration (core)"],
            ["auth/", "Authentication plugins (core + oauth2)"],
            ["course/", "Core course view — already patched for public browse"],
            ["enrol/", "Enrolment plugins; fee + self used"],
            ["local/", "IIIDEM local plugins"],
            ["mod/", "Activities (quiz, assign, customcert, zoom, webexactivity, …)"],
            ["payment/", "Payment accounts; custom Razorpay/PNB/ICICI"],
            ["register/", "Custom registration UI"],
            ["theme/iiidem2/", "IIIDEM theme — primary customisation surface"],
            ["theme/boost/", "Parent theme (core)"],
            ["user/", "Core user pages"],
            ["moodledata/", "Dataroot — outside public serving in production"],
            ["docs/", "KT, security, and ops markdown"],
            ["local_dev_logs/", "Developer CLI probes — not for production"],
            [".ddev/", "Local containers, nginx snippets, Dex"],
        ],
    )

    add_heading(doc, "15. People and roles (typical)", 1)
    table(
        doc,
        ["Role", "What they do on this LMS"],
        [
            ["Student", "Register, pay (certificate course), learn, upload IMW readings, take quizzes, get certificate"],
            ["Teacher / editing teacher", "Course content, live class, attendance, Meet your Professors cards"],
            ["Manager / admin", "Users, enrolments, payments, theme settings, upgrades"],
            ["Guest", "Browse public course marketing pages only"],
        ],
    )
    para(
        doc,
        "Site admins are excluded from Meet your Professors (e.g. Jaindra). Do not use "
        "a site-admin account as a featured professor.",
    )

    add_heading(doc, "16. How to take a change from idea to production", 1)
    numbered(
        doc,
        [
            "Reproduce on DDEV (course 4 vs course 8).",
            "Change only theme_iiidem2 or an existing IIIDEM plugin.",
            "Bump theme version; add upgrade.php savepoint; purge caches locally.",
            "Verify in the browser: guest, logged-in student, teacher, the other course.",
            "Commit when asked. Do not commit secrets or moodledata.",
            "Deploy files to staging, upgrade with upgrade key, retest.",
            "Deploy to production the same way. Confirm upgrade key screen is expected, not an error.",
        ],
    )

    add_heading(doc, "17. Documents to read next", 1)
    bullets(
        doc,
        [
            "docs/site-flow-diagrams.md — registration, payment, live class, backup flows",
            "docs/restrict-admin-access.md — upgrade key and admin IP allowlist",
            "docs/file-upload-security.md — why uploads are locked down",
            "docs/LOCAL-SERVER-TESTING.md — sharing local DDEV on LAN",
            "docs/daily-backup-cron.md — staging backup paths",
            "theme/iiidem2/settings.php — all theme admin settings",
        ],
    )

    add_heading(doc, "18. Handover contacts and secrets (fill in)", 1)
    para(
        doc,
        "Do not put live passwords in this file. Record who holds them.",
    )
    table(
        doc,
        ["Item", "Held by / where"],
        [
            ["Production SSH / panel", ""],
            ["MOODLE_UPGRADEKEY", "Server env or config.production.php"],
            ["MariaDB production", "config.production.php"],
            ["Razorpay keys", "Moodle payment account / paygw settings"],
            ["Webex OAuth", "local_iiidem_webexattendance admin settings"],
            ["Dex / SSO (if enabled)", ".ddev/dex and IdP consoles"],
            ["Git remote", "Ask current maintainer"],
            ["CDAC audit evidence", "docs/ plus scanner attachments"],
        ],
    )

    add_heading(doc, "19. One-page summary for the incoming person", 1)
    para(
        doc,
        "You are inheriting a Moodle 4.5 site with a heavy custom theme (iiidem2), not a "
        "greenfield app. Course 4 is the paid certificate programme. Course 7 on production "
        "(course 8 locally) is the IMW workshop with a shared reading library. Never lower "
        "the theme version. Never copy local config.php onto production. After uploading "
        "theme files, complete the upgrade (upgrade key on the web UI). Custom behaviour "
        "almost always lives in theme/iiidem2/. If a guest page fatals, check that every "
        "new PHP class was uploaded and the upgrade ran.",
    )

    add_heading(doc, "Document control", 1)
    table(
        doc,
        ["Field", "Value"],
        [
            ["Title", "IIIDEM LMS Knowledge Transfer"],
            ["System", "Moodle 4.5.11+ / theme_iiidem2 2024101120"],
            ["Prepared", date.today().isoformat()],
            ["Classification", "Internal — IIIDEM / ECI"],
            ["Format", "Microsoft Word (.docx)"],
        ],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print("Wrote", OUT)


if __name__ == "__main__":
    build()
